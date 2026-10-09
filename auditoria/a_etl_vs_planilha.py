"""AUDITORIA (a): planilha -> dados.json.

Recalcula, direto do .xlsx e sem usar gerar_dados.py, os campos de cada
registro e as entregas de cada empreendimento, e compara com dados.json.
"""
import json, re, sys, unicodedata
import openpyxl

XLSX = sys.argv[1]
DADOS = sys.argv[2]

TIPOS = ["Grandes e Médias", "Governadores", "Migrado Novo PAC"]
SIT_ETAPA = {"Contratada": "contratado", "Contratação parcial": "contratado",
             "Em licitação": "contratado", "Em andamento": "contratado",
             "Concluído": "contratado", "Em ação preparatória": "aContratar",
             "Desistência": "desistencia", "Cancelada": "desistencia",
             "Habilitada": "habilitada",
             "Desistência da Habilitação": "desistenciaHabilitacao"}
IND = {"corredorKm": 13, "trilhosKm": 14, "viarioKm": 15, "cicloKm": 16,
       "estacoes": 17, "terminais": 18, "viadutos": 19, "passarelas": 20,
       "pontes": 21, "abrigos": 22, "outrasOAE": 23, "veiculos": 24,
       "its": 25, "cco": 26, "patios": 27, "projetos": 28, "extProjetadaKm": 29}

def txt(v):
    if v is None: return None
    s = str(v).strip()
    return s or None

def num(v):
    if isinstance(v, (int, float)): return float(v)
    if v is None: return 0.0
    s = str(v).replace("R$", "").replace(".", "").replace(",", ".").strip()
    try: return float(s)
    except ValueError: return 0.0

erros = []
avisos = []


def check(cond, msg):
    if not cond: erros.append(msg)

wb = openpyxl.load_workbook(XLSX, data_only=True)
base = wb["BASEDEDADOS"]
cab = {}
for c in range(1, base.max_column + 1):
    t = txt(base.cell(2, c).value)
    if t and t not in cab: cab[t] = c

def col(nome):
    for k, v in cab.items():
        if k.strip().lower() == nome.strip().lower(): return v
    return None

C = {k: col(v) for k, v in {
    "tipo": "Tipo", "uf": "UF", "municipio": "Município Principal",
    "situacao": "Situação Contrato", "apoio": "Apoio",
    "valorContratado": "Valor contratado", "empreendimento": "Empreendimento",
    "unid": "Unid", "qnt": "Qnt",
}.items()}
faltando = [k for k, v in C.items() if not v]
check(not faltando, f"colunas nao encontradas na BASEDEDADOS: {faltando}")

pacote = json.load(open(DADOS, encoding="utf-8"))
regs = {r["id"]: r for r in pacote["registros"]}

# ---- 1. campos de cada registro -------------------------------------------
vistos = 0
NOME_DA_UF = {
    "AC": "Acre", "AL": "Alagoas", "AM": "Amazonas", "AP": "Amapá", "BA": "Bahia",
    "CE": "Ceará", "DF": "Distrito Federal", "ES": "Espírito Santo", "GO": "Goiás",
    "MA": "Maranhão", "MG": "Minas Gerais", "MS": "Mato Grosso do Sul",
    "MT": "Mato Grosso", "PA": "Pará", "PB": "Paraíba", "PE": "Pernambuco",
    "PI": "Piauí", "PR": "Paraná", "RJ": "Rio de Janeiro",
    "RN": "Rio Grande do Norte", "RO": "Rondônia", "RR": "Roraima",
    "RS": "Rio Grande do Sul", "SC": "Santa Catarina", "SE": "Sergipe",
    "SP": "São Paulo", "TO": "Tocantins",
}


def sem_estado(municipio, uf):
    if not municipio:
        return municipio
    sufixo = ", " + NOME_DA_UF.get((uf or "").strip().upper(), "\0")
    return municipio[: -len(sufixo)].strip() if municipio.endswith(sufixo) else municipio


for l in range(3, base.max_row + 1):
    tipo = txt(base.cell(l, C["tipo"]).value)
    if not tipo: continue
    r = regs.get(l)
    if r is None:
        check(False, f"linha {l} da BASEDEDADOS ausente em dados.json")
        continue
    vistos += 1
    for campo, coluna in (("tipo", "tipo"), ("uf", "uf"), ("municipio", "municipio"),
                          ("situacao", "situacao"), ("empreendimento", "empreendimento")):
        esperado = txt(base.cell(l, C[coluna]).value)
        if campo == "municipio":
            # O ETL tira o nome do estado colado no do município
            # ("Natal, Rio Grande do Norte" -> "Natal"), senão a mesma cidade
            # conta duas vezes. A auditoria aplica a mesma regra ao esperado.
            esperado = sem_estado(esperado, txt(base.cell(l, C["uf"]).value))
        check(r[campo] == esperado,
              f"linha {l}: {campo} planilha={esperado!r} json={r[campo]!r}")
    for campo in ("apoio", "valorContratado"):
        esperado = round(num(base.cell(l, C[campo]).value), 2)
        check(abs(r[campo] - esperado) < 0.01,
              f"linha {l}: {campo} planilha={esperado} json={r[campo]}")
    check(r["noEscopo"] == (tipo in TIPOS), f"linha {l}: noEscopo incoerente")
    sit = txt(base.cell(l, C["situacao"]).value)
    check(r["etapa"] == SIT_ETAPA.get(sit), f"linha {l}: etapa de {sit!r}")
    check(r["migrado"] == (tipo == "Migrado Novo PAC"), f"linha {l}: migrado")
    # km/unidades derivam de Quantidade + Unidade de Medida
    unid = (txt(base.cell(l, C["unid"]).value) or "").lower()
    q = num(base.cell(l, C["qnt"]).value)
    km = q if unid.startswith("km") else 0.0
    un = q if unid.startswith("unid") else 0.0
    check(abs(r["km"] - round(km, 3)) < 0.01, f"linha {l}: km {km} != {r['km']}")
    check(abs(r["unidades"] - round(un, 3)) < 0.01, f"linha {l}: unidades {un} != {r['unidades']}")

check(vistos == len(regs), f"registros: planilha={vistos} json={len(regs)}")

# ---- 2. entregas por empreendimento ---------------------------------------
ws = wb["Indicadores de Obra"]
# A aba cresce a cada levantamento novo: o fim dos dados é a linha de TOTAL
# (ou a última preenchida), nunca um número fixo — com 136 cravado, as linhas
# acrescentadas depois passavam despercebidas pela auditoria.
COL_EMP_IND = next(c for c in range(1, ws.max_column + 1)
                   if str(ws.cell(5, c).value or "").startswith("Empreendimento"))
FIM_IND = 6
for _l in range(6, ws.max_row + 2):
    _v = str(ws.cell(_l, COL_EMP_IND).value or "").strip()
    if not _v or _v.upper().startswith("TOTAL"):
        FIM_IND = _l
        break
ancoras = {r["id"]: r for r in pacote["registros"] if r.get("entregas")}
linhas_aba = 0
for l in range(6, FIM_IND):
    ids = [int(n) for n in re.findall(r"\d+", str(ws.cell(l, 3).value or ""))]
    if not ids: continue
    linhas_aba += 1
    anc = [i for i in ids if i in ancoras]
    check(len(anc) == 1,
          f"aba Indicadores linha {l}: {len(anc)} ancora(s) para as linhas {ids}")
    if len(anc) != 1: continue
    e = ancoras[anc[0]]["entregas"]
    for chave, c in IND.items():
        esperado = round(num(ws.cell(l, c).value), 3)
        check(abs(e[chave] - esperado) < 0.001,
              f"linha {l} ({chave}): aba={esperado} json={e[chave]}")
    check(abs(e["oae"] - sum(round(num(ws.cell(l, IND[k]).value), 3)
                             for k in ("viadutos", "pontes", "outrasOAE"))) < 0.001,
          f"linha {l}: oae != viadutos+pontes+outras")
    # tipologia: igual em todos os registros do empreendimento
    tip = txt(ws.cell(l, 34).value)
    for i in ids:
        if i not in regs:
            continue
        if not tip:
            # Empreendimento novo, ainda sem a coluna "Tipologia do
            # empreendimento" preenchida: o ETL usa a reserva por modo e
            # categoria. Vale registrar para a equipe preencher, mas não é
            # divergência de cálculo.
            avisos.append(f"linha {l}: sem tipologia na aba; o painel usou a "
                          f"reserva '{regs[i].get('tipologia')}' "
                          f"({(regs[i].get('empreendimento') or '')[:40]})")
            continue
        check(regs[i].get("tipologia") == tip,
              f"linha {l}: tipologia do registro {i} = {regs[i].get('tipologia')!r}, aba={tip!r}")

check(linhas_aba == len(ancoras),
      f"empreendimentos: aba={linhas_aba} ancoras no json={len(ancoras)}")

# ---- 3. cobertura: todo registro no escopo pertence a um empreendimento ----
no_escopo = [r for r in pacote["registros"]
             if r["noEscopo"] and r["etapa"] in ("contratado", "aContratar", "desistencia")]
ids_aba = set()
for l in range(6, FIM_IND):
    ids_aba |= {int(n) for n in re.findall(r"\d+", str(ws.cell(l, 3).value or ""))}
fora = [r["id"] for r in no_escopo if r["id"] not in ids_aba]
check(not fora, f"registros no escopo sem empreendimento na aba: {fora}")

print(f"registros conferidos: {vistos} · empreendimentos: {linhas_aba}")
if avisos:
    print(f"AVISOS ({len(avisos)}) — não são divergências de cálculo:")
    for a in avisos:
        print("  -", a)
print(f"ERROS: {len(erros)}")
for e in erros[:40]: print("  -", e)
sys.exit(1 if erros else 0)
