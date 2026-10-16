#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
================================================================================
 GERAR_DADOS.PY — Camada de dados do Painel Grandes e Médias Cidades (Novo PAC)
 Ministério das Cidades / Secretaria Nacional de Mobilidade — SEMOB
================================================================================

O QUE ESTE SCRIPT FAZ
---------------------
Lê a aba analítica da planilha de acompanhamento (uma linha por proposta) e
produz a camada de dados do painel:

    Excel (aba BASEDEDADOS  +  aba Indicadores de Obra)
        |
        v
    gerar_dados.py          <-- mapeamento de colunas + regras de negócio
        |
        +--> dados.json                       (consumo via servidor HTTP)
        +--> assets/js/dados-embutido.js      (consumo via file:// , duplo clique)
        |
        v
    index.html

PRINCÍPIO ADOTADO
-----------------
O script NÃO pré-calcula os indicadores do painel. Ele entrega a base
normalizada (registros) + metadados. Toda agregação é feita no navegador,
em função dos filtros ativos. Consequência prática: incluir um novo filtro,
um novo gráfico ou um novo recorte NÃO exige alterar este script.

ONDE MEXER QUANDO A PLANILHA MUDAR
----------------------------------
Somente no bloco "MAPEAMENTO DA PLANILHA" abaixo (nome da aba, linha de
cabeçalho e nomes das colunas). O restante do arquivo é regra de negócio
estável.

COMO USAR
---------
    pip install openpyxl
    python gerar_dados.py
    python gerar_dados.py --planilha dados/minha_planilha.xlsx
    python gerar_dados.py --mock          # gera base fictícia (sem Excel)

================================================================================
"""

import argparse
import json
import os
import random
import re
import sys
import unicodedata
from datetime import datetime

try:
    from zoneinfo import ZoneInfo
    FUSO = ZoneInfo("America/Sao_Paulo")
except ImportError:  # Python < 3.9
    from datetime import timezone, timedelta
    FUSO = timezone(timedelta(hours=-3))


# =============================================================================
#  1. MAPEAMENTO DA PLANILHA  —  ÚNICO BLOCO A AJUSTAR QUANDO A BASE MUDAR
# =============================================================================

PLANILHA_PADRAO = "dados/teste_grandes_e_medias.xlsx"

ABA_BASE = "BASEDEDADOS"     # aba analítica: uma linha = uma proposta
LINHA_CABECALHO = 2          # linha onde estão os títulos das colunas
LINHA_PRIMEIRO_DADO = 3      # primeira linha de dados

# Chave interna do painel  ->  título da coluna na planilha.
# Os títulos são comparados de forma tolerante (sem acento, sem caixa,
# sem espaços extras), então pequenas variações de digitação não quebram.
COLUNAS = {
    "proposta":        "Proposta",
    "anoProposta":     "Ano Proposta",
    "contrato":        "Contrato",
    "tipo":            "Tipo",
    "uf":              "UF",
    "regiao":          "Região",
    "proponente":      "Proponente",
    "tipoProponente":  "Tipo de Proponente",
    "municipio":       "Município Principal",
    "agente":          "Agente Financeiro",
    "programa":        "Programa",
    "empreendimento":  "Empreendimento",
    "fonte":           "Fonte",
    "apoio":           "Apoio",
    "contrapartida":   "Contrapartida",
    "situacao":        "Situação Contrato",
    "dataAssinatura":  "Data assinatura",
    "execucao":        "Situação da execução",
    "valorContratado": "Valor contratado",
    "valorLiberado":   "Valor liberado/desembolsado",
    "observacoes":     "Observações",
    "portaria":        "Portaria",
    "anoPortaria":     "Ano portaria",
    "item":            "Item principal ",
    "unid":            "Unid",
    "qnt":             "Qnt",
    "categoria":       "Categoria do item",
    "modo":            "Modo de transporte",
}

# Colunas obrigatórias: sem elas o painel não se sustenta.
COLUNAS_OBRIGATORIAS = ["tipo", "uf", "situacao", "apoio", "valorContratado"]


# -----------------------------------------------------------------------------
#  Aba "Indicadores de Obra": tudo o que cada empreendimento ENTREGA
#  (corredores, trilhos, estações, terminais, OAE, abrigos, veículos...).
#  Uma linha da aba = um empreendimento, que pode reunir várias linhas da
#  BASEDEDADOS (coluna "Linha(s) na BASEDEDADOS"). As quantidades entram no
#  painel UMA única vez por empreendimento, ancoradas em um único registro.
# -----------------------------------------------------------------------------
ABA_INDICADORES = "Indicadores de Obra"
LINHA_CABECALHO_IND = 5
LINHA_PRIMEIRO_DADO_IND = 6

# Tipologia do empreendimento: classificação única (e exclusiva) por
# empreendimento, preenchida na aba 'Indicadores de Obra'. Alimenta a aba
# "Investimento" e o filtro "Tipo de investimento" do painel.
TITULO_TIPOLOGIA = "Tipologia do empreendimento"
TIPOLOGIAS = ["Metrô e trens", "VLT", "BRT e corredores de ônibus",
              "Sistema viário e OAE", "Mobilidade ativa",
              "Terminais e sistemas", "Estudos e projetos"]
# Plano B para registros sem tipologia preenchida (a partir do modo/categoria).
MODO_PARA_TIPOLOGIA = {
    "Metrô": TIPOLOGIAS[0], "Trens": TIPOLOGIAS[0], "Aeromóvel": TIPOLOGIAS[0],
    "VLTs": TIPOLOGIAS[1], "BRTs": TIPOLOGIAS[2], "Corredores de Ônibus": TIPOLOGIAS[2],
    "Sistema Viário": TIPOLOGIAS[3], "OAE (Obra de Arte Especial)": TIPOLOGIAS[3],
    "Ciclovias": TIPOLOGIAS[4], "Terminais": TIPOLOGIAS[5], "Sistemas": TIPOLOGIAS[5],
    "Abrigos": TIPOLOGIAS[5],
}


def tipologia_reserva(reg):
    if reg.get("categoria") == "Estudos e Projetos":
        return TIPOLOGIAS[6]
    return MODO_PARA_TIPOLOGIA.get(reg.get("modo"), TIPOLOGIAS[5])


IND_COLUNAS_BASE = {
    "linhas":       "Linha(s) na BASEDEDADOS",
    "modalidade":   "Modalidade",
    "levantamento": "Situação do levantamento",
    "idGoverna":    "ID Governa",
}

# chave interna -> (título da coluna, rótulo no painel, unidade, grupo)
# A ordem desta lista é a ordem de exibição no painel.
INDICADORES = [
    ("corredorKm",     "Corredor / BRT / faixa exclusiva (km)", "Corredores, BRT e faixas exclusivas", "km",   "Vias e sistemas"),
    ("trilhosKm",      "Trilhos: metrô, VLT, trem (km)",        "Trilhos (metrô, VLT e trem)",         "km",   "Vias e sistemas"),
    ("viarioKm",       "Sistema Viário (km)",                   "Sistema viário",                      "km",   "Vias e sistemas"),
    ("cicloKm",        "Ciclovia / ciclofaixa (km)",            "Ciclovias e ciclofaixas",             "km",   "Vias e sistemas"),
    ("estacoes",       "Estações (un.)",                        "Estações",                            "un.",  "Pontos de parada e integração"),
    ("terminais",      "Terminais (un.)",                       "Terminais",                           "un.",  "Pontos de parada e integração"),
    ("abrigos",        "Abrigos / paradas (un.)",               "Abrigos e paradas",                   "un.",  "Pontos de parada e integração"),
    ("oae",            None,                                    "Obras de arte especiais (OAE)",       "un.",  "Obras de arte"),
    ("passarelas",     "Passarelas (un.)",                      "Passarelas",                          "un.",  "Obras de arte"),
    ("veiculos",       "Veículos / material rodante (un.)",     "Veículos e material rodante",         "un.",  "Sistemas e material rodante"),
    ("its",            "ITS/Semáforos inteligentes (un.)",      "ITS e semáforos inteligentes",        "un.",  "Sistemas e material rodante"),
    ("cco",            "CCO (un.)",                             "Centros de controle (CCO)",           "un.",  "Sistemas e material rodante"),
    ("patios",         "Pátios / oficinas / garagens (un.)",    "Pátios, oficinas e garagens",         "un.",  "Sistemas e material rodante"),
    ("projetos",       "Projetos / estudos (un.)",              "Projetos e estudos",                  "un.",  "Estudos e projetos"),
    ("extProjetadaKm", "Extensão projetada (km)",               "Extensão projetada em estudos",       "km",   "Estudos e projetos"),
]

# Colunas lidas diretamente da aba (a OAE é a soma das três abaixo).
IND_COLUNAS_VALOR = {c[0]: c[1] for c in INDICADORES if c[1]}
IND_COLUNAS_VALOR.update({
    "viadutos":  "Viadutos (un.)",
    "pontes":    "Pontes (un.)",
    "outrasOAE": "Outras OAE: túneis, trincheiras, elevados (un.)",
})
# Decisão de 30/09/2026: viadutos + pontes + outras OAE formam UM indicador
# (algumas propostas só informam o agregado, ex.: "pontes e viadutos").
# Passarelas permanecem separadas.
COMPONENTES_OAE = ["viadutos", "pontes", "outrasOAE"]


# =============================================================================
#  2. REGRAS DE NEGÓCIO  —  estáveis, revisar apenas em mudança de conceito
# =============================================================================

# Recorte do painel. "Avançar Cidades" existe na base, mas é outro programa:
# fica fora do consolidado (conferido contra a aba "Dados para Painel").
TIPOS_NO_ESCOPO = ["Grandes e Médias", "Governadores", "Migrado Novo PAC"]

# Situação do contrato -> etapa do funil.
SITUACAO_PARA_ETAPA = {
    "Contratada":                 "contratado",
    "Contratação parcial":        "contratado",
    "Em licitação":               "contratado",
    "Em andamento":               "contratado",
    "Concluído":                  "contratado",
    "Em ação preparatória":       "aContratar",
    "Desistência":                "desistencia",
    "Cancelada":                  "desistencia",
    "Habilitada":                 "habilitada",
    "Desistência da Habilitação": "desistenciaHabilitacao",
}

# Etapas que compõem o universo "selecionado" (a carteira viva + baixas).
ETAPAS_SELECIONADAS = ["contratado", "aContratar", "desistencia"]

# Modos cuja métrica natural é extensão (km). Usados no controle de qualidade.
MODOS_LINEARES = ["Metrô", "VLTs", "Trens", "BRTs", "Corredores de Ônibus",
                  "Ciclovias", "Aeromóvel"]

# Classificação definida pela SEMOB para itens sem categoria/modo na planilha:
# item (minúsculo) -> (categoria do item, modo de transporte).
CLASSIFICACAO_DE_ITEM = {
    "semáforos inteligentes":    ("Equipamentos", "Sistemas"),   # Sorocaba
    "trem metropolitano":        ("Obras",        "Trens"),      # Linhas 7 e 9, São Paulo
    "trilhos - material rodante": ("Equipamentos", "Trens"),     # Recife
}

# Complemento ao nome do empreendimento (definido pela SEMOB, 16/10/2026): o
# levantamento de Recife consolida 4 linhas da base e registra 6 veículos (os
# 6 TUEs do Metrô de BH), que o nome da linha principal não mencionava.
COMPLEMENTO_DO_NOME = {
    "Recuperação da infraestrutura do Metrô do Recife - substituição de trilhos e dormentes":
        " (inclui aquisição de 6 TUEs do Metrô de BH)",
}

# Alertas já conferidos pela SEMOB (16/10/2026): a situação é a esperada e o
# registro não deve mais aparecer em "Qualidade da base". Chave: (nº da
# proposta, alerta). Se o dado mudar a ponto de o alerta deixar de existir, a
# entrada simplesmente não tem efeito; se mudar para outro problema, ele aparece.
ALERTAS_CONFERIDOS = {
    ("56000002550/2023", "contratadaSemValor"),   # Vila Velha
    ("56000002277/2023", "contratadaSemValor"),   # São José dos Pinhais
    ("40974880",         "contratadaSemValor"),   # Brasília (Migrado)
    ("680974",           "contratadaSemValor"),   # Campinas (Migrado)
    ("26744/2024",       "unidadeSuspeita"),      # Salvador, telecom. VLT
    ("56000004011/2023", "unidadeSuspeita"),      # Salvador, telecom. VLT
}

# Equipamentos medidos em unidades (material rodante) não são "unidade suspeita"
# num modo linear: a unidade é o próprio veículo.
ITENS_EM_UNIDADES = {"trilhos - material rodante"}

# Rótulo aplicado a propostas anteriores ao Novo PAC (sem ano de portaria).
ROTULO_SEM_ANO = "Anterior a 2023"


# =============================================================================
#  3. TEXTOS INSTITUCIONAIS  —  editar aqui, não no HTML
# =============================================================================

TEXTOS = {
    "tituloPainel": "Painel de Monitoramento — Grandes e Médias Cidades",
    "programa": "Novo PAC · Mobilidade Urbana Sustentável",
    "subtitulo": "Infraestrutura de transporte público de média e alta capacidade",
    "orgao": "Secretaria Nacional de Mobilidade — SEMOB",
    "resumo": [
        "O subeixo Mobilidade Grandes e Médias Cidades do Novo PAC financia a "
        "infraestrutura estruturante do transporte público urbano: metrôs, trens "
        "urbanos, VLTs, BRTs e corredores de ônibus, além de terminais, estações, "
        "sistemas embarcados, requalificação viária e mobilidade ativa. Os recursos "
        "combinam Orçamento Geral da União e financiamento (FGTS/Pró-Transporte e "
        "FAT), operados pela CAIXA e pelo BNDES.",

        "A carteira acompanhada percorre um fluxo em três etapas — seleção, "
        "contratação e execução. O painel monitora a conversão entre essas "
        "etapas e, principalmente, o que cada real contratado entrega em "
        "infraestrutura: quilômetros de via de metrô, VLT, trem, BRT e corredor, "
        "além das unidades de material rodante, terminais, abrigos e sistemas.",
    ],
    "fonte": "Elaboração: Ministério das Cidades / SEMOB — Coordenação de Monitoramento. "
             "Base: planilha de acompanhamento de propostas do Novo PAC Mobilidade.",
    "aviso": "Documento de trabalho para uso interno. Valores sujeitos a "
             "revisão conforme atualização da base.",
}


# =============================================================================
#  4. UTILITÁRIOS
# =============================================================================

def log(nivel, msg):
    prefixo = {"ok": "[OK]   ", "info": "[INFO] ", "erro": "[ERRO] ",
               "aviso": "[AVISO]"}
    print(f"  {prefixo.get(nivel, '')} {msg}")


def normalizar(texto):
    """Compara títulos de coluna ignorando acento, caixa e espaços."""
    if texto is None:
        return ""
    t = unicodedata.normalize("NFKD", str(texto))
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", t).strip().lower()


def texto(valor):
    """Limpa strings da planilha (espaços supérfluos, zero-width, vazios)."""
    if valor is None:
        return None
    if isinstance(valor, str):
        v = valor.replace("\u200b", "").replace("\xa0", " ").strip()
        return v or None
    return valor


def numero(valor):
    """Converte para float tolerando texto e vírgula decimal."""
    if isinstance(valor, (int, float)):
        return float(valor)
    if isinstance(valor, str):
        v = valor.replace("\u200b", "").replace(".", "").replace(",", ".").strip()
        try:
            return float(v)
        except ValueError:
            return 0.0
    return 0.0


NOME_DA_UF = {
    "AC": "Acre", "AL": "Alagoas", "AM": "Amazonas", "AP": "Amapá", "BA": "Bahia",
    "CE": "Ceará", "DF": "Distrito Federal", "ES": "Espírito Santo", "GO": "Goiás",
    "MA": "Maranhão", "MG": "Minas Gerais", "MS": "Mato Grosso do Sul",
    "MT": "Mato Grosso", "PA": "Pará", "PB": "Paraíba", "PE": "Pernambuco",
    "PI": "Piauí", "PR": "Paraná", "RJ": "Rio de Janeiro",
    "RN": "Rio Grande do Norte", "RO": "Rondônia", "RR": "Roraima",
    "RS": "Rio Grande do Sul", "SC": "Santa Catarina", "SE": "Sergipe",
    "SP": "São Paulo", "TO": "Tocantins",
}


def municipio_limpo(nome, uf):
    """Tira o nome do estado quando ele vem colado no do município.

    A planilha ora traz "Natal", ora "Natal, Rio Grande do Norte" — a mesma
    cidade contada duas vezes no cartão "Alcance territorial". O sufixo é
    redundante (a UF já está na coluna ao lado), então sai aqui, no único
    ponto de leitura.
    """
    if not nome:
        return nome
    sufixo = ", " + NOME_DA_UF.get((uf or "").strip().upper(), "\0")
    return nome[: -len(sufixo)].strip() if nome.endswith(sufixo) else nome


def ano_portaria(valor):
    """Normaliza o ano da portaria.

    Aceita 2024, '2024', '2025/2026' (vale o ano mais recente) e vazio.
    Retorna (ano_int_ou_None, rotulo_para_o_eixo).
    """
    if valor is None or valor == "":
        return None, ROTULO_SEM_ANO
    anos = re.findall(r"(19|20)\d{2}", str(valor))
    if not anos:
        return None, ROTULO_SEM_ANO
    achados = [int(a) for a in re.findall(r"(?:19|20)\d{2}", str(valor))]
    ano = max(achados)
    return ano, str(ano)


def iso(valor):
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    return None


# =============================================================================
#  5. LEITURA DA PLANILHA
# =============================================================================

def localizar_colunas(ws):
    """Casa os títulos da planilha com o dicionário COLUNAS."""
    titulos = {}
    for c in range(1, ws.max_column + 1):
        t = normalizar(ws.cell(LINHA_CABECALHO, c).value)
        if t and t not in titulos:
            titulos[t] = c

    indices, faltando = {}, []
    for chave, titulo in COLUNAS.items():
        col = titulos.get(normalizar(titulo))
        if col:
            indices[chave] = col
        else:
            faltando.append(f"{chave} ('{titulo}')")

    if faltando:
        log("aviso", f"{len(faltando)} coluna(s) não localizada(s): "
                     f"{', '.join(faltando)}")
    criticas = [c for c in COLUNAS_OBRIGATORIAS if c not in indices]
    if criticas:
        raise SystemExit(
            f"Colunas obrigatórias ausentes: {', '.join(criticas)}. "
            f"Ajuste o bloco MAPEAMENTO DA PLANILHA em gerar_dados.py."
        )
    return indices


def ler_registros(caminho):
    import openpyxl

    if not os.path.exists(caminho):
        raise SystemExit(f"Planilha não encontrada: {caminho}")

    wb = openpyxl.load_workbook(caminho, data_only=True)
    if ABA_BASE not in wb.sheetnames:
        raise SystemExit(f"Aba '{ABA_BASE}' não encontrada. "
                         f"Abas disponíveis: {', '.join(wb.sheetnames)}")
    ws = wb[ABA_BASE]
    idx = localizar_colunas(ws)

    def celula(linha, chave):
        col = idx.get(chave)
        return ws.cell(linha, col).value if col else None

    registros = []
    for r in range(LINHA_PRIMEIRO_DADO, ws.max_row + 1):
        tipo = texto(celula(r, "tipo"))
        uf = texto(celula(r, "uf"))
        if not tipo and not uf:
            continue  # linha vazia / rodapé

        situacao = texto(celula(r, "situacao"))
        etapa = SITUACAO_PARA_ETAPA.get(situacao)

        unid = (texto(celula(r, "unid")) or "").lower()
        qnt = numero(celula(r, "qnt"))
        ano, rotulo = ano_portaria(celula(r, "anoPortaria"))

        reg = {
            "id": r,
            "proposta": texto(celula(r, "proposta")),
            "anoProposta": texto(celula(r, "anoProposta")),
            "contrato": texto(celula(r, "contrato")),
            "tipo": tipo,
            "programa": texto(celula(r, "programa")),
            "uf": uf,
            "regiao": texto(celula(r, "regiao")),
            "municipio": municipio_limpo(texto(celula(r, "municipio")), uf),
            "proponente": texto(celula(r, "proponente")),
            "tipoProponente": texto(celula(r, "tipoProponente")),
            "agente": texto(celula(r, "agente")),
            "fonte": texto(celula(r, "fonte")),
            "empreendimento": texto(celula(r, "empreendimento")),
            "categoria": texto(celula(r, "categoria")),
            "modo": texto(celula(r, "modo")),
            "item": texto(celula(r, "item")),
            "situacao": situacao,
            "etapa": etapa,
            "execucao": texto(celula(r, "execucao")),
            "apoio": numero(celula(r, "apoio")),
            "contrapartida": numero(celula(r, "contrapartida")),
            "valorContratado": numero(celula(r, "valorContratado")),
            "valorLiberado": numero(celula(r, "valorLiberado")),
            "km": qnt if unid == "km" else 0.0,
            "unidades": qnt if unid.startswith("unid") else 0.0,
            "unidadeOriginal": texto(celula(r, "unid")),
            "quantidade": qnt,
            "anoPortaria": ano,
            "rotuloAno": rotulo,
            "portaria": texto(celula(r, "portaria")),
            "dataAssinatura": iso(celula(r, "dataAssinatura")),
            "observacoes": texto(celula(r, "observacoes")),
            "noEscopo": tipo in TIPOS_NO_ESCOPO,
            # Migrado Novo PAC: "apoio" é o recurso federal ORIGINAL do
            # empreendimento; "valorContratado" é só a parcela migrada.
            "migrado": tipo == "Migrado Novo PAC",
            # Preenchido por anexar_entregas() em UM registro por empreendimento.
            "entregas": None,
        }
        # Categoria nova "Sistema Viário" (06/10/2026): se as tabelas de apoio
        # da planilha ainda não a conhecerem, o painel a reconhece aqui.
        if (reg["item"] or "").strip().lower() == "sistema viário":
            reg["categoria"] = reg["categoria"] if reg["categoria"] not in (None, "Não classificado") else "Obras"
            reg["modo"] = reg["modo"] or "Sistema Viário"

        # Item que começa por "Projeto" é, por definição, estudo e projeto
        # (09/10/2026). As colunas de apoio "Categoria do item" e "Modo de
        # transporte" vieram vazias para 34 desses registros, e com isso a
        # fatia "Estudos e Projetos" sumia de "Natureza do apoio" e virava
        # "Não classificado". Aqui o próprio nome do item responde: não é
        # palpite, é leitura literal do dado.
        if (reg["item"] or "").strip().lower().startswith("projeto"):
            if reg["categoria"] in (None, "", "Não classificado"):
                reg["categoria"] = "Estudos e Projetos"
            reg["modo"] = reg["modo"] or "Estudos e Projetos"
        # Itens que a planilha (16/10/2026) ainda traz sem categoria e sem modo.
        # A classificação abaixo foi definida pela SEMOB: vale até as tabelas de
        # apoio da planilha reconhecerem esses itens; depois, a planilha manda.
        _cls = CLASSIFICACAO_DE_ITEM.get((reg["item"] or "").strip().lower())
        if _cls:
            if reg["categoria"] in (None, "", "Não classificado"):
                reg["categoria"] = _cls[0]
            reg["modo"] = reg["modo"] or _cls[1]
        if reg["empreendimento"] in COMPLEMENTO_DO_NOME:
            reg["empreendimento"] += COMPLEMENTO_DO_NOME[reg["empreendimento"]]
        reg["alertas"] = auditar(reg)
        _ok = [a for a in reg["alertas"] if (str(reg["proposta"]).strip(), a) in ALERTAS_CONFERIDOS]
        if _ok:
            reg["alertas"] = [a for a in reg["alertas"] if a not in _ok]
            reg["alertasConferidos"] = _ok
        registros.append(reg)

    entregas_meta = anexar_entregas(wb, registros)
    for r_ in registros:
        if not r_.get("tipologia"):
            r_["tipologia"] = tipologia_reserva(r_)
    return registros, entregas_meta



# =============================================================================
#  5.1 INDICADORES DE OBRA  (aba "Indicadores de Obra")
# =============================================================================

def _linhas_da_base(valor):
    """'10, 73, 513, 514' -> [10, 73, 513, 514]"""
    return [int(n) for n in re.findall(r"\d+", str(valor or ""))]


def _escolher_ancora(candidatos, modalidade):
    """Escolhe, entre as linhas da BASEDEDADOS de um empreendimento, a que
    carrega as quantidades. Prefere a mesma modalidade da aba, depois a que
    está contratada (maior valor), para que os filtros do painel (modalidade,
    visão contratado/selecionado, UF) vejam o empreendimento onde ele vive."""
    alvo = normalizar(modalidade)
    mesma = [r for r in candidatos if normalizar(r["tipo"]) == alvo] or candidatos
    viva = [r for r in mesma if r["etapa"] in ETAPAS_SELECIONADAS] or mesma
    return max(viva, key=lambda r: (r["etapa"] == "contratado",
                                    r["valorContratado"], r["apoio"]))


def anexar_entregas(wb, registros):
    """Lê a aba 'Indicadores de Obra' e grava as quantidades em `entregas`
    de UM registro por empreendimento (o 'âncora'). Devolve metadados."""
    if ABA_INDICADORES not in wb.sheetnames:
        log("aviso", f"Aba '{ABA_INDICADORES}' não encontrada — o painel "
                     f"seguirá apenas com o item principal de cada proposta.")
        return None

    ws = wb[ABA_INDICADORES]
    titulos = {}
    for c in range(1, ws.max_column + 1):
        t = normalizar(ws.cell(LINHA_CABECALHO_IND, c).value)
        if t and t not in titulos:
            titulos[t] = c

    def col(titulo):
        return titulos.get(normalizar(titulo))

    base = {k: col(t) for k, t in IND_COLUNAS_BASE.items()}
    valores = {k: col(t) for k, t in IND_COLUNAS_VALOR.items()}
    faltando = [t for k, t in list(IND_COLUNAS_BASE.items()) + list(IND_COLUNAS_VALOR.items())
                if not col(t) and k != "idGoverna"]
    if faltando:
        raise SystemExit("Colunas ausentes na aba 'Indicadores de Obra': "
                         + "; ".join(faltando))

    col_tipologia = col(TITULO_TIPOLOGIA)
    if not col_tipologia:
        log("aviso", f"Coluna '{TITULO_TIPOLOGIA}' ausente — será usada a "
                     f"classificação de reserva (modo/categoria).")
    por_id = {r["id"]: r for r in registros}
    ancoras, vistos = [], set()
    status = {}
    problemas = []

    for linha in range(LINHA_PRIMEIRO_DADO_IND, ws.max_row + 1):
        rotulo = texto(ws.cell(linha, base["modalidade"]).value) if base["modalidade"] else None
        celula_f = texto(ws.cell(linha, 6).value)
        if isinstance(celula_f, str) and celula_f.upper().startswith("TOTAL"):
            continue
        linhas_base = _linhas_da_base(ws.cell(linha, base["linhas"]).value)
        if not linhas_base:
            continue

        candidatos = [por_id[n] for n in linhas_base if n in por_id]
        if not candidatos:
            problemas.append(f"linha {linha}: BASEDEDADOS {linhas_base} inexistente")
            continue

        ancora = _escolher_ancora(candidatos, rotulo)
        if ancora["id"] in vistos:
            problemas.append(f"linha {linha}: registro {ancora['id']} já ancora outro empreendimento")
            continue
        vistos.add(ancora["id"])

        if col_tipologia:
            tip = texto(ws.cell(linha, col_tipologia).value)
            if tip and tip not in TIPOLOGIAS:
                problemas.append(f"linha {linha}: tipologia '{tip}' fora da lista prevista")
                tip = None
            if tip:
                for c_ in candidatos:
                    c_["tipologia"] = tip

        q = {k: numero(ws.cell(linha, c).value) for k, c in valores.items()}
        q["oae"] = sum(q[k] for k in COMPONENTES_OAE)
        ancora["entregas"] = {k: round(v, 3) for k, v in q.items()}
        ancora["idGoverna"] = texto(ws.cell(linha, base["idGoverna"]).value) if base["idGoverna"] else None

        lev = texto(ws.cell(linha, base["levantamento"]).value) or "Não informado"
        status[lev] = status.get(lev, 0) + 1
        if lev not in ("Completo", "Não se aplica") and ancora["alertas"] is not None:
            ancora["alertas"].append("levantamentoIncompleto")

        # Conferência de consistência: tipo/UF dos registros agrupados.
        ufs = {r["uf"] for r in candidatos}
        if len(ufs) > 1:
            problemas.append(f"linha {linha}: linhas da base em UFs diferentes {sorted(ufs)}")
        ancoras.append(ancora)

    for p in problemas:
        log("aviso", p)

    totais = {k: round(sum(a["entregas"][k] for a in ancoras), 2)
              for k in IND_COLUNAS_VALOR.keys() | {"oae"}}
    log("ok", f"{len(ancoras)} empreendimento(s) com indicadores de obra "
              f"(situação do levantamento: {status}).")

    return {
        "aba": ABA_INDICADORES,
        "empreendimentos": len(ancoras),
        "situacaoLevantamento": status,
        "totais": totais,
        "dicionario": [
            {"chave": c, "rotulo": r, "unidade": u, "grupo": g}
            for c, _, r, u, g in INDICADORES
        ],
        "composicaoOAE": COMPONENTES_OAE,
        "regras": [
            "Cada empreendimento entra uma única vez, ancorado em um registro da BASEDEDADOS.",
            "Célula vazia em empreendimento com levantamento 'Completo' vale zero.",
            "OAE = viadutos + pontes + outras OAE (túneis, trincheiras, elevados); passarelas ficam à parte.",
        ],
    }


# =============================================================================
#  6. CONTROLE DE QUALIDADE DA BASE
# =============================================================================

ALERTAS = {
    "semModo":            "Sem modo de transporte classificado",
    "semUnidade":         "Quantidade informada sem unidade de medida",
    "unidadeSuspeita":    "Modo linear medido em 'unid' (verificar se seria km)",
    "contratadaSemValor": "Contratada sem valor contratado registrado",
    "valorAcimaApoio":    "Valor contratado maior que o apoio selecionado",
    "semAnoPortaria":     "Sem ano de portaria",
    "situacaoDesconhecida": "Situação do contrato fora da lista prevista",
    "levantamentoIncompleto": "Indicadores de obra ainda não levantados por completo",
}


def auditar(reg):
    """Marca inconsistências que merecem conferência antes de divulgar."""
    a = []
    if not reg["situacao"] or reg["etapa"] is None:
        a.append("situacaoDesconhecida")
    if not reg["modo"]:
        a.append("semModo")
    if reg["quantidade"] and not reg["unidadeOriginal"]:
        a.append("semUnidade")
    if reg["modo"] in MODOS_LINEARES and reg["unidades"] > 0 and \
            (reg["item"] or "").strip().lower() not in ITENS_EM_UNIDADES:
        a.append("unidadeSuspeita")
    if reg["etapa"] == "contratado" and reg["valorContratado"] <= 0:
        a.append("contratadaSemValor")
    if reg["valorContratado"] > reg["apoio"] + 0.01:
        a.append("valorAcimaApoio")
    if reg["anoPortaria"] is None and reg["tipo"] != "Migrado Novo PAC":
        a.append("semAnoPortaria")
    return a


# =============================================================================
#  7. METADADOS E CONFERÊNCIA
# =============================================================================

def ordenar_unicos(valores):
    return sorted({v for v in valores if v})


def montar_meta(registros, origem, data_str, entregas=None):
    escopo = [r for r in registros if r["noEscopo"]]
    contratados = [r for r in escopo if r["etapa"] == "contratado"]
    selecionados = [r for r in escopo if r["etapa"] in ETAPAS_SELECIONADAS]

    # A auditoria reportada considera apenas o recorte do painel: registros de
    # outros programas ficam fora e não devem inflar o indicador de qualidade.
    contagem_alertas = {}
    for r in escopo:
        for a in r["alertas"]:
            contagem_alertas[a] = contagem_alertas.get(a, 0) + 1

    anos = sorted({r["anoPortaria"] for r in escopo if r["anoPortaria"]})

    return {
        "dataAtualizacao": data_str,
        "geradoEm": datetime.now(FUSO).isoformat(timespec="seconds"),
        "origem": os.path.basename(origem),
        "escopo": {
            "tiposNoEscopo": TIPOS_NO_ESCOPO,
            "registrosTotais": len(registros),
            "registrosNoEscopo": len(escopo),
        },
        "conferencia": {
            "propostasContratadas": len(contratados),
            "valorContratado": round(sum(r["valorContratado"] for r in contratados), 2),
            "propostasSelecionadas": len(selecionados),
            # Mesma regra do painel (Regras.valorDe): Migrado Novo PAC usa o valor
            # contratado como "selecionado"; os demais, o apoio.
            "valorSelecionado": round(sum(
                r["valorContratado"] if r["tipo"] == "Migrado Novo PAC" else r["apoio"]
                for r in selecionados), 2),
        },
        "opcoes": {
            "cenarios": ["Consolidado"] + TIPOS_NO_ESCOPO,
            "anos": [ROTULO_SEM_ANO] + [str(a) for a in anos],
            "regioes": ordenar_unicos(r["regiao"] for r in escopo),
            "ufs": ordenar_unicos(r["uf"] for r in escopo),
            "modos": ordenar_unicos(r["modo"] for r in escopo),
            "tipologias": [t for t in TIPOLOGIAS if any(r.get("tipologia") == t for r in escopo)],
            "categorias": ordenar_unicos(r["categoria"] for r in escopo),
            "fontes": ordenar_unicos(r["fonte"] for r in escopo),
            "situacoes": ordenar_unicos(r["situacao"] for r in escopo),
        },
        "entregas": entregas,
        "dicionarioAlertas": ALERTAS,
        "qualidade": {
            "registrosComAlerta": sum(1 for r in escopo if r["alertas"]),
            "porTipo": contagem_alertas,
        },
        "textos": TEXTOS,
    }


def conferir(meta):
    """Imprime os totais de fechamento para conferência contra a planilha."""
    c = meta["conferencia"]
    print("-" * 72)
    log("info", "Conferência (comparar com a aba 'Dados para Painel'):")
    def brl(v):
        return ("R$ " + f"{v:,.2f}").replace(",", "@").replace(".", ",").replace("@", ".")

    print(f"           Propostas contratadas ....... {c['propostasContratadas']}")
    print(f"           Valor contratado ............ {brl(c['valorContratado'])}")
    print(f"           Propostas selecionadas ...... {c['propostasSelecionadas']}")
    print(f"           Valor selecionado ........... {brl(c['valorSelecionado'])}")
    q = meta["qualidade"]
    if q["registrosComAlerta"]:
        log("aviso", f"{q['registrosComAlerta']} registro(s) com alerta de "
                     f"consistência — ver seção 'Qualidade da base' no painel.")
        for chave, n in sorted(q["porTipo"].items(), key=lambda kv: -kv[1]):
            print(f"           · {n:>3}  {ALERTAS.get(chave, chave)}")


# =============================================================================
#  8. BASE FICTÍCIA (MOCK) — usada só com --mock, nunca em produção
# =============================================================================

def gerar_mock(quantidade=90):
    """Base sintética para desenvolver a interface sem a planilha real.

    ATENÇÃO: dados INVENTADOS. O JSON gerado sai marcado com "mock": true e o
    painel exibe uma tarja de aviso quando detecta essa marca.
    """
    random.seed(42)
    ufs = {"SP": "Sudeste", "RJ": "Sudeste", "MG": "Sudeste", "BA": "Nordeste",
           "CE": "Nordeste", "PE": "Nordeste", "PR": "Sul", "RS": "Sul",
           "DF": "Centro-Oeste", "GO": "Centro-Oeste", "PA": "Norte"}
    modos = ["Metrô", "VLTs", "Trens", "BRTs", "Corredores de Ônibus",
             "Terminais", "Sistemas", "Ciclovias"]
    situacoes = ["Contratada"] * 5 + ["Em ação preparatória"] * 2 + \
                ["Habilitada"] * 2 + ["Cancelada"]

    registros = []
    for i in range(quantidade):
        uf = random.choice(list(ufs))
        modo = random.choice(modos)
        situacao = random.choice(situacoes)
        etapa = SITUACAO_PARA_ETAPA[situacao]
        apoio = round(random.uniform(8e6, 1.5e9), 2)
        linear = modo in MODOS_LINEARES
        reg = {
            "id": 1000 + i,
            "proposta": f"MOCK-{i:04d}",
            "anoProposta": random.choice([2023, 2024, 2025, 2026]),
            "contrato": None, "tipo": random.choice(TIPOS_NO_ESCOPO),
            "programa": "MOCK", "uf": uf, "regiao": ufs[uf],
            "municipio": f"Município {i:02d}", "proponente": f"Proponente {i:02d}",
            "tipoProponente": random.choice(["Município", "Estado"]),
            "agente": random.choice(["CAIXA", "BNDES"]),
            "fonte": random.choice(["OGU", "FIN"]),
            "empreendimento": f"[MOCK] Intervenção de {modo} — exemplo {i:02d}",
            "categoria": random.choice(["Obras", "Equipamentos", "Estudos e Projetos"]),
            "modo": modo, "item": modo, "situacao": situacao, "etapa": etapa,
            "execucao": None, "apoio": apoio, "contrapartida": 0.0,
            "valorContratado": round(apoio * random.uniform(0.6, 1.0), 2)
                               if etapa == "contratado" else 0.0,
            "valorLiberado": 0.0,
            "km": round(random.uniform(1, 30), 2) if linear else 0.0,
            "unidades": 0.0 if linear else float(random.randint(1, 40)),
            "unidadeOriginal": "km" if linear else "unid",
            "quantidade": 0.0,
            "observacoes": None, "portaria": None, "dataAssinatura": None,
            "noEscopo": True,
        }
        ano = random.choice([None, 2023, 2024, 2025, 2026])
        reg["anoPortaria"] = ano
        reg["rotuloAno"] = str(ano) if ano else ROTULO_SEM_ANO
        reg["quantidade"] = reg["km"] or reg["unidades"]
        reg["alertas"] = auditar(reg)
        registros.append(reg)
    return registros


# =============================================================================
#  9. ESCRITA DOS ARQUIVOS
# =============================================================================

def escrever(pacote, caminho_json, caminho_js):
    with open(caminho_json, "w", encoding="utf-8") as f:
        json.dump(pacote, f, ensure_ascii=False, indent=1)

    os.makedirs(os.path.dirname(caminho_js), exist_ok=True)
    compacto = json.dumps(pacote, ensure_ascii=False, separators=(",", ":"))
    with open(caminho_js, "w", encoding="utf-8") as f:
        f.write("/* Gerado por gerar_dados.py — não editar à mão.\n"
                "   Espelha dados.json para permitir abrir o painel por duplo\n"
                "   clique (file://), onde o fetch() é bloqueado pelo navegador. */\n")
        f.write("window.DADOS_PAINEL = " + compacto + ";\n")


def main():
    p = argparse.ArgumentParser(
        description="Gera a camada de dados do Painel Grandes e Médias Cidades.")
    p.add_argument("--planilha", default=PLANILHA_PADRAO)
    p.add_argument("--saida-json", default="dados.json")
    p.add_argument("--saida-js", default="assets/js/dados-embutido.js")
    p.add_argument("--mock", action="store_true",
                   help="gera base fictícia, sem ler a planilha")
    args = p.parse_args()

    agora = datetime.now(FUSO)
    print("=" * 72)
    print(" PAINEL GRANDES E MÉDIAS CIDADES — geração da base de dados")
    print(f" {agora.strftime('%d/%m/%Y %H:%M')}")
    print("=" * 72)

    if args.mock:
        log("aviso", "Modo MOCK: os dados gerados são fictícios.")
        registros = gerar_mock()
        entregas = None
        origem = "MOCK"
    else:
        log("info", f"Lendo '{ABA_BASE}' de {args.planilha}")
        registros, entregas = ler_registros(args.planilha)
        origem = args.planilha

    log("ok", f"{len(registros)} registro(s) lido(s).")

    meta = montar_meta(registros, origem, agora.strftime("%d/%m/%Y"), entregas)
    meta["mock"] = bool(args.mock)
    conferir(meta)

    escrever({"meta": meta, "registros": registros},
             args.saida_json, args.saida_js)

    print("-" * 72)
    log("ok", f"Gerado: {args.saida_json}")
    log("ok", f"Gerado: {args.saida_js}")
    print("=" * 72)


if __name__ == "__main__":
    main()
